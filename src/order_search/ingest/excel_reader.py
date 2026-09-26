from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .fields import parse_date_range, split_barcodes, to_int

SHEET_NAME = "Order Data"
INT_COLUMNS = ("Dept", "Class", "Subclass", "NO.")
DATE_COLUMN = "Original Expected Date"
BARCODE_COLUMN = "Bar_Code"


class SourceFormatError(RuntimeError):
    """Raised when the workbook does not look like the expected export."""


def _dedupe_headers(values: list[Any]) -> list[str]:
    """Excel allows duplicate and blank header cells; give each a usable name."""
    seen: dict[str, int] = {}
    out: list[str] = []
    for index, raw in enumerate(values):
        name = "" if raw is None else str(raw).strip()
        if not name:
            name = f"column_{index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def read_order_rows(path: str | Path) -> list[dict[str, Any]]:
    """Reads the 'Order Data' worksheet into a list of dicts keyed by header.

    'Bar_Code' becomes a list of barcodes, 'Original Expected Date' becomes
    'date_from' and 'date_to', and the numeric columns become ints.
    """
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise SourceFormatError(
                f"worksheet {SHEET_NAME!r} not found;"
                f" found {sorted(workbook.sheetnames)}"
            )
        sheet = workbook[SHEET_NAME]
        stream = sheet.iter_rows(values_only=True)
        try:
            header = next(stream)
        except StopIteration:
            raise SourceFormatError(
                f"worksheet {SHEET_NAME!r} is empty"
            ) from None

        headers = _dedupe_headers(list(header))
        rows: list[dict[str, Any]] = []
        for values in stream:
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            record: dict[str, Any] = dict(zip(headers, values))
            record[BARCODE_COLUMN] = split_barcodes(record.get(BARCODE_COLUMN))
            start, end = parse_date_range(record.get(DATE_COLUMN))
            record["date_from"] = start
            record["date_to"] = end
            for column in INT_COLUMNS:
                if column in record:
                    record[column] = to_int(record[column])
            rows.append(record)
        return rows
    finally:
        workbook.close()
