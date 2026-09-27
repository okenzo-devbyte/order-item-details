import base64
import json
import os
import subprocess
import sys
from pathlib import Path

from api.snapshot import SnapshotStore

from conftest import DATA_FILE

ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "sample_order_data_1000_records.xlsx"
PYTHON = ROOT / ".venv" / "Scripts" / "python.exe"


def run(args, env_extra=None):
    env = dict(os.environ)
    env["PYTHONPATH"] = str(ROOT / "src")
    env["PYTHONIOENCODING"] = "utf-8"
    if env_extra:
        env.update(env_extra)
    return subprocess.run(
        [str(PYTHON), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=ROOT,
        env=env,
    )


class TestImportScript:
    def test_builds_a_database_and_prints_the_report(self, tmp_path):
        target = tmp_path / "snapshot.db"
        result = run(
            [
                "scripts/import_excel.py",
                str(DATA_FILE),
                "--out",
                str(target),
            ]
        )
        assert result.returncode == 0, result.stderr
        for label in [
            "rows read",
            "rows imported",
            "products",
            "customers",
            "barcodes",
        ]:
            assert label in result.stdout
        assert target.is_file()

    def test_reports_127_barcodes(self, tmp_path):
        target = tmp_path / "snapshot.db"
        result = run(
            ["scripts/import_excel.py", str(DATA_FILE), "--out", str(target)]
        )
        assert "barcodes       : 127" in result.stdout

    def test_missing_file_exits_with_a_message(self, tmp_path):
        result = run(
            [
                "scripts/import_excel.py",
                str(tmp_path / "nope.xlsx"),
                "--out",
                str(tmp_path / "x.db"),
            ]
        )
        assert result.returncode != 0
        assert "nope.xlsx" in result.stderr


class TestFixtureScript:
    def test_generates_a_workbook_with_varied_dates(self, tmp_path):
        target = tmp_path / "fixture.xlsx"
        result = run(["scripts/make_fixture.py", "--out", str(target)])
        assert result.returncode == 0, result.stderr
        assert target.is_file()

        from openpyxl import load_workbook

        sheet = load_workbook(target)["Order Data"]
        values = {
            row[12] for row in sheet.iter_rows(min_row=2, values_only=True)
        }
        assert len(values) > 5, "fixture must contain several distinct dates"


class TestQueryScript:
    def test_searches_a_customer(self, tmp_path):
        db = tmp_path / "snapshot.db"
        assert run(
            ["scripts/import_excel.py", str(DATA_FILE), "--out", str(db)]
        ).returncode == 0
        result = run(["scripts/query_cli.py", "--db", str(db), "สุรชัย"])
        assert result.returncode == 0, result.stderr
        payload = json.loads(result.stdout)
        assert payload["mode"] == "customer"
        assert payload["customers"][0]["order_count"] == 65

    def test_searches_a_barcode(self, tmp_path):
        db = tmp_path / "snapshot.db"
        run(["scripts/import_excel.py", str(DATA_FILE), "--out", str(db)])
        result = run(
            ["scripts/query_cli.py", "--db", str(db), "8850250001234"]
        )
        payload = json.loads(result.stdout)
        assert payload["detected"] == "barcode"
        assert payload["products"][0]["name"] == "น้ำดื่มสิงห์ 600 มล. x 12"

    def test_accepts_filters(self, tmp_path):
        db = tmp_path / "snapshot.db"
        run(["scripts/import_excel.py", str(DATA_FILE), "--out", str(db)])
        result = run(
            [
                "scripts/query_cli.py",
                "--db",
                str(db),
                "--store",
                "101",
                "--order-type",
                "Pickup",
                "น้ำดื่มสิงห์ 600 มล. x 12",
            ]
        )
        payload = json.loads(result.stdout)
        assert payload["products"][0]["order_count"] < 86

    def test_suggest_endpoint(self, tmp_path):
        db = tmp_path / "snapshot.db"
        run(["scripts/import_excel.py", str(DATA_FILE), "--out", str(db)])
        result = run(
            ["scripts/query_cli.py", "--db", str(db), "--suggest", "น้ำ"]
        )
        payload = json.loads(result.stdout)
        assert len(payload["suggestions"]) == 4


RAW_KEY = bytes(range(32))
KEY = base64.urlsafe_b64encode(RAW_KEY).decode("ascii")


def run_script(script, args, env_extra=None):
    env = dict(os.environ)
    env["PYTHONPATH"] = str(ROOT / "src")
    env["PYTHONIOENCODING"] = "utf-8"
    env["DATA_KEY"] = KEY
    env["DATA_KEY_ID"] = "v1"
    if env_extra:
        env.update(env_extra)
    return subprocess.run(
        [str(PYTHON), str(ROOT / "scripts" / script)] + args,
        capture_output=True,
        text=True,
        encoding="utf-8",
        cwd=ROOT,
        env=env,
    )


def test_build_snapshot_produces_a_searchable_file(tmp_path):
    out = tmp_path / "snapshot.enc"
    result = run_script(
        "build_snapshot.py", [str(DATA_FILE), "--out", str(out)]
    )
    assert result.returncode == 0, result.stderr
    assert "products" in result.stdout
    store = SnapshotStore({"v1": RAW_KEY}, "v1")
    store.load_file(out)
    assert len(store.search("น้ำ")["products"]) == 4


def test_make_admin_creates_a_login(tmp_path):
    db_path = tmp_path / "app.db"
    result = run_script(
        "make_admin.py",
        ["--db", str(db_path), "--username", "boss", "--password", "password123"],
    )
    assert result.returncode == 0, result.stderr
    from api.app_db import AppDB
    from api.users import authenticate

    db = AppDB(db_path)
    assert authenticate(db, "boss", "password123")["role"] == "admin"
    db.close()
